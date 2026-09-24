import os
from datetime import timedelta
import docx
from django.core.management.base import BaseCommand
from django.contrib.auth import get_user_model
from django.utils import timezone
from django.core.files import File
from accounts.models import CandidateProfile, RecruiterProfile
from jobs.models import Job, SavedJob
from applications.models import Application
from interviews.models import Interview
from notifications.models import Notification
from resume_matching.matcher import calculate_resume_match
from resume_matching.parser import extract_text_from_file

User = get_user_model()

class Command(BaseCommand):
    help = 'Seeds initial realistic demo data for HireFlow platform.'

    def handle(self, *args, **options):
        self.stdout.write(self.style.NOTICE('Starting HireFlow demo data seeding...'))

        # 1. Create Admin
        admin_email = 'admin@hireflow.com'
        admin_user, created = User.objects.get_or_create(
            email=admin_email,
            defaults={
                'name': 'Alex Rivera (Platform Admin)',
                'role': User.Role.ADMIN,
                'phone': '+1 (555) 019-2831',
                'is_staff': True,
                'is_superuser': True,
                'is_active': True,
            }
        )
        if created:
            admin_user.set_password('Admin@123')
            admin_user.save()
            self.stdout.write(self.style.SUCCESS(f"Created Admin: {admin_email} / Admin@123"))
        else:
            self.stdout.write(f"Admin {admin_email} already exists.")

        # 2. Create Recruiters
        # Recruiter 1: TechCorp Innovations
        r1_email = 'recruiter@techcorp.com'
        r1, created = User.objects.get_or_create(
            email=r1_email,
            defaults={
                'name': 'Sarah Jenkins',
                'role': User.Role.RECRUITER,
                'phone': '+1 (555) 234-5678',
                'is_active': True,
            }
        )
        if created:
            r1.set_password('Recruiter@123')
            r1.save()
            RecruiterProfile.objects.create(
                user=r1,
                company_name='TechCorp Innovations',
                company_description='Leading cloud software and distributed systems provider powering next-gen enterprise tools.',
                website='https://techcorp.example.com',
                location='San Francisco, CA',
                industry='Software & Cloud Infrastructure',
                company_size='500-1000'
            )
            self.stdout.write(self.style.SUCCESS(f"Created Recruiter 1: {r1_email} / Recruiter@123"))

        # Recruiter 2: FinFlow Solutions
        r2_email = 'recruiter@finflow.io'
        r2, created = User.objects.get_or_create(
            email=r2_email,
            defaults={
                'name': 'Marcus Vance',
                'role': User.Role.RECRUITER,
                'phone': '+1 (555) 987-6543',
                'is_active': True,
            }
        )
        if created:
            r2.set_password('Recruiter@123')
            r2.save()
            RecruiterProfile.objects.create(
                user=r2,
                company_name='FinFlow Solutions',
                company_description='Next-generation financial data aggregation and real-time transaction processing API platform.',
                website='https://finflow.example.com',
                location='New York, NY',
                industry='Fintech & Banking Technology',
                company_size='100-250'
            )
            self.stdout.write(self.style.SUCCESS(f"Created Recruiter 2: {r2_email} / Recruiter@123"))

        # 3. Create Sample Resume Files
        media_resumes_dir = os.path.join('media', 'resumes')
        os.makedirs(media_resumes_dir, exist_ok=True)

        sample_resume_1_path = os.path.join(media_resumes_dir, 'sample_resume_greeshma.docx')
        if not os.path.exists(sample_resume_1_path):
            doc = docx.Document()
            doc.add_heading('Greeshma Rani - Full Stack Engineer', 0)
            doc.add_paragraph('Email: candidate@hireflow.com | Location: Austin, TX | GitHub: github.com/greeshma')
            doc.add_heading('Summary', level=1)
            doc.add_paragraph('Experienced full stack developer with expertise in building responsive web applications using Python, Django, REST API, React, SQL, and Docker.')
            doc.add_heading('Skills', level=1)
            doc.add_paragraph('Languages & Frameworks: Python, Django, JavaScript, React, SQL, MySQL, PostgreSQL, REST API, Git, Docker, HTML, CSS.')
            doc.add_heading('Experience', level=1)
            doc.add_paragraph('Software Engineer at CloudCraft (2023 - Present): Designed REST APIs in Django, built interactive dashboards in React, managed MySQL databases.')
            doc.add_heading('Education', level=1)
            doc.add_paragraph('Bachelor of Technology in Computer Science (2020 - 2024)')
            doc.save(sample_resume_1_path)

        # 4. Create Candidates
        c1_email = 'candidate@hireflow.com'
        c1, created = User.objects.get_or_create(
            email=c1_email,
            defaults={
                'name': 'Greeshma Rani',
                'role': User.Role.CANDIDATE,
                'phone': '+1 (555) 456-7890',
                'is_active': True,
            }
        )
        if created:
            c1.set_password('Candidate@123')
            c1.save()
            profile = CandidateProfile.objects.create(
                user=c1,
                location='Austin, TX',
                bio='Dedicated Full-Stack Developer passionate about clean code, scalable APIs, and intuitive user experiences.',
                education='Bachelor of Technology',
                degree='B.Tech in Computer Science',
                institution='State University of Technology',
                graduation_year=2024,
                experience=3,
                skills='Python, Django, React, SQL, Git, REST API, Docker, MySQL',
                projects='HireFlow Recruitment Platform, Real-time Chat App, Automated Testing Framework'
            )
            with open(sample_resume_1_path, 'rb') as f:
                profile.resume.save('sample_resume_greeshma.docx', File(f), save=True)
            self.stdout.write(self.style.SUCCESS(f"Created Candidate 1: {c1_email} / Candidate@123"))

        c2_email = 'candidate2@hireflow.com'
        c2, created = User.objects.get_or_create(
            email=c2_email,
            defaults={
                'name': 'David Chen',
                'role': User.Role.CANDIDATE,
                'phone': '+1 (555) 789-0123',
                'is_active': True,
            }
        )
        if created:
            c2.set_password('Candidate@123')
            c2.save()
            CandidateProfile.objects.create(
                user=c2,
                location='Seattle, WA',
                bio='Backend specialist focusing on Node.js, TypeScript, microservices, and cloud architectures.',
                education='Master of Science',
                degree='M.S. in Software Engineering',
                institution='Pacific Institute of Technology',
                graduation_year=2022,
                experience=4,
                skills='JavaScript, TypeScript, React, Node.js, Express.js, MongoDB, AWS, Docker',
                projects='High-throughput payment gateway, Microservices telemetry orchestrator'
            )
            self.stdout.write(self.style.SUCCESS(f"Created Candidate 2: {c2_email} / Candidate@123"))

        # 5. Create Jobs
        today = timezone.now().date()
        job1, created = Job.objects.get_or_create(
            recruiter=r1,
            title='Senior Python / Django Developer',
            defaults={
                'company': 'TechCorp Innovations',
                'department': 'Backend Engineering',
                'description': 'We are looking for an experienced Senior Python/Django developer to design high-throughput microservices, optimize MySQL queries, and build robust RESTful APIs.',
                'responsibilities': '- Architect and maintain Django REST APIs\n- Collaborate with frontend engineers to build seamless integrations\n- Optimize MySQL queries and database schema\n- Implement unit and integration test suites',
                'qualifications': '- 3+ years experience with Python & Django\n- Solid understanding of relational databases (MySQL/PostgreSQL)\n- Proficient in Git, Docker, and REST APIs',
                'location': 'San Francisco, CA (Hybrid)',
                'job_type': Job.JobType.FULL_TIME,
                'experience_min': 2,
                'experience_max': 6,
                'salary_min': 110000.00,
                'salary_max': 140000.00,
                'required_skills': 'Python, Django, SQL, REST API, Git, Docker',
                'openings': 2,
                'deadline': today + timedelta(days=30),
                'status': Job.Status.ACTIVE,
            }
        )

        job2, created = Job.objects.get_or_create(
            recruiter=r1,
            title='Full Stack React & Node Engineer',
            defaults={
                'company': 'TechCorp Innovations',
                'department': 'Product Development',
                'description': 'Join our product team to deliver modern, responsive SaaS dashboards and full stack capabilities utilizing React and Node.js.',
                'responsibilities': '- Build responsive web applications using React\n- Design backend services with Node.js and Express\n- Ensure high code quality and test coverage',
                'qualifications': '- Strong proficiency in JavaScript/TypeScript and React\n- Experience building RESTful APIs with Node.js\n- Familiarity with CI/CD and Git',
                'location': 'Remote (US)',
                'job_type': Job.JobType.REMOTE,
                'experience_min': 2,
                'experience_max': 5,
                'salary_min': 95000.00,
                'salary_max': 125000.00,
                'required_skills': 'React, JavaScript, TypeScript, Node.js, Express.js, REST API, Git',
                'openings': 3,
                'deadline': today + timedelta(days=25),
                'status': Job.Status.ACTIVE,
            }
        )

        job3, created = Job.objects.get_or_create(
            recruiter=r2,
            title='Backend API Architect',
            defaults={
                'company': 'FinFlow Solutions',
                'department': 'Core Infrastructure',
                'description': 'Seeking an experienced Backend Architect to lead real-time transactional systems and API security for high-volume financial data streams.',
                'responsibilities': '- Drive architecture for core API gateways\n- Ensure sub-millisecond query responses\n- Enforce robust authentication and audit standards',
                'qualifications': '- Deep proficiency in Python (FastAPI/Django) or Go\n- Experience with Redis, PostgreSQL, and Docker/Kubernetes',
                'location': 'New York, NY',
                'job_type': Job.JobType.FULL_TIME,
                'experience_min': 4,
                'experience_max': 8,
                'salary_min': 135000.00,
                'salary_max': 170000.00,
                'required_skills': 'Python, FastAPI, SQL, PostgreSQL, Redis, Docker, Kubernetes, Microservices',
                'openings': 1,
                'deadline': today + timedelta(days=45),
                'status': Job.Status.ACTIVE,
            }
        )

        job4, created = Job.objects.get_or_create(
            recruiter=r2,
            title='Frontend UI/UX Specialist',
            defaults={
                'company': 'FinFlow Solutions',
                'department': 'Design & Frontend',
                'description': 'Create sleek, accessible, responsive client dashboards for our financial analytics engine.',
                'responsibilities': '- Translate Figma prototypes into responsive web code\n- Implement smooth micro-interactions\n- Maintain component library',
                'qualifications': '- 2+ years of modern React and CSS experience\n- Strong eye for typography, layout, and performance',
                'location': 'New York, NY (Hybrid)',
                'job_type': Job.JobType.FULL_TIME,
                'experience_min': 2,
                'experience_max': 4,
                'salary_min': 85000.00,
                'salary_max': 110000.00,
                'required_skills': 'React, CSS, HTML, JavaScript, TypeScript, Next.js',
                'openings': 1,
                'deadline': today + timedelta(days=20),
                'status': Job.Status.PENDING_APPROVAL,
            }
        )

        job5, created = Job.objects.get_or_create(
            recruiter=r1,
            title='Cloud Infrastructure & DevOps Engineer',
            defaults={
                'company': 'TechCorp Innovations',
                'department': 'DevOps & SRE',
                'description': 'Maintain CI/CD pipelines, Kubernetes clusters, and automated monitoring.',
                'responsibilities': '- Manage container orchestration and deployments\n- Automate test/build/deploy pipelines',
                'qualifications': '- Strong Linux and bash scripting\n- Hands-on Docker, Kubernetes, and AWS experience',
                'location': 'Remote',
                'job_type': Job.JobType.CONTRACT,
                'experience_min': 3,
                'experience_max': 7,
                'salary_min': 105000.00,
                'salary_max': 135000.00,
                'required_skills': 'Linux, Docker, Kubernetes, AWS, CI/CD, Terraform, Git',
                'openings': 1,
                'deadline': today + timedelta(days=15),
                'status': Job.Status.DRAFT,
            }
        )
        self.stdout.write(self.style.SUCCESS("Created demo Jobs in various statuses."))

        # 6. Saved Jobs
        SavedJob.objects.get_or_create(candidate=c1, job=job1)
        SavedJob.objects.get_or_create(candidate=c1, job=job3)

        # 7. Create Applications
        app1, created = Application.objects.get_or_create(
            candidate=c1,
            job=job1,
            defaults={
                'cover_letter': 'I am thrilled to apply for the Senior Python / Django Developer role. My hands-on experience building production Django APIs and managing MySQL schemas makes me a strong fit for your team.',
                'additional_information': 'Available to start within two weeks. Open to hybrid arrangements in Austin/SF.',
                'status': Application.Status.SHORTLISTED,
                'match_score': 83,
            }
        )
        if created and c1.candidate_profile.resume:
            app1.resume = c1.candidate_profile.resume
            app1.save()

        app2, created = Application.objects.get_or_create(
            candidate=c1,
            job=job3,
            defaults={
                'cover_letter': 'Excited about the Backend API Architect opening at FinFlow Solutions. I have extensive experience with Python, SQL, and microservices.',
                'additional_information': 'Portfolio and GitHub links available on my profile.',
                'status': Application.Status.APPLIED,
                'match_score': 62,
            }
        )
        if created and c1.candidate_profile.resume:
            app2.resume = c1.candidate_profile.resume
            app2.save()

        app3, created = Application.objects.get_or_create(
            candidate=c2,
            job=job2,
            defaults={
                'cover_letter': 'Dear Hiring Team, I specialize in full stack React and Node.js solutions and would love to contribute to TechCorp Innovations.',
                'additional_information': 'Authorized to work in the US without sponsorship.',
                'status': Application.Status.INTERVIEW_SCHEDULED,
                'match_score': 86,
            }
        )
        if created and c1.candidate_profile.resume:
            app3.resume = c1.candidate_profile.resume
            app3.save()

        self.stdout.write(self.style.SUCCESS("Created demo Applications."))

        # 8. Create Interview
        tomorrow = today + timedelta(days=1)
        interview1, created = Interview.objects.get_or_create(
            application=app1,
            candidate=c1,
            recruiter=r1,
            job=job1,
            defaults={
                'date': tomorrow,
                'time': '14:00:00',
                'interview_type': Interview.InterviewType.TECHNICAL,
                'meeting_link': 'https://meet.google.com/hireflow-tech-interview-demo',
                'interviewer': 'Sarah Jenkins & Tech Lead',
                'notes': 'Discussion on Django ORM performance, database indexing in MySQL, and architecture of distributed services.',
                'status': Interview.Status.SCHEDULED,
            }
        )
        self.stdout.write(self.style.SUCCESS("Created demo Interview."))

        # 9. Create Notifications
        Notification.objects.get_or_create(
            recipient=c1,
            title="Interview Scheduled: Senior Python / Django Developer",
            defaults={
                'message': f"Your Technical interview has been scheduled for tomorrow at 2:00 PM with Sarah Jenkins.",
                'notification_type': Notification.NotificationType.INTERVIEW,
                'is_read': False,
            }
        )

        Notification.objects.get_or_create(
            recipient=c1,
            title="Application Shortlisted!",
            defaults={
                'message': f"TechCorp Innovations shortlisted your application for 'Senior Python / Django Developer'!",
                'notification_type': Notification.NotificationType.APPLICATION,
                'is_read': True,
            }
        )

        Notification.objects.get_or_create(
            recipient=r1,
            title="New Applicant: Greeshma Rani",
            defaults={
                'message': "Greeshma Rani applied for 'Senior Python / Django Developer' with an 83% Resume Match Score.",
                'notification_type': Notification.NotificationType.APPLICATION,
                'is_read': False,
            }
        )

        Notification.objects.get_or_create(
            recipient=admin_user,
            title="Job Pending Approval: Frontend UI/UX Specialist",
            defaults={
                'message': "Marcus Vance (FinFlow Solutions) submitted a new job for admin review.",
                'notification_type': Notification.NotificationType.JOB,
                'is_read': False,
            }
        )
        self.stdout.write(self.style.SUCCESS("Created demo Notifications."))

        self.stdout.write(self.style.SUCCESS("""
=====================================================
HireFlow Demo Data Successfully Seeded!
-----------------------------------------------------
Admin Demo:
  Email:    admin@hireflow.com
  Password: Admin@123

Recruiter Demo 1:
  Email:    recruiter@techcorp.com
  Password: Recruiter@123

Recruiter Demo 2:
  Email:    recruiter@finflow.io
  Password: Recruiter@123

Candidate Demo 1:
  Email:    candidate@hireflow.com
  Password: Candidate@123

Candidate Demo 2:
  Email:    candidate2@hireflow.com
  Password: Candidate@123
=====================================================
"""))
