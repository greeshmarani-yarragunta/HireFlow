import io
import os
import pypdf
import docx

def extract_text_from_file(file_obj_or_path):
    """
    Extracts plain text from either a file path or an in-memory/uploaded file object.
    Supports PDF and DOCX.
    """
    filename = getattr(file_obj_or_path, 'name', '') or str(file_obj_or_path)
    ext = os.path.splitext(filename)[1].lower()

    text = ""
    try:
        if ext == '.pdf':
            # Handle both file path and uploaded File / BytesIO
            if isinstance(file_obj_or_path, (str, os.PathLike)):
                with open(file_obj_or_path, 'rb') as f:
                    reader = pypdf.PdfReader(f)
                    for page in reader.pages:
                        extracted = page.extract_text()
                        if extracted:
                            text += extracted + "\n"
            else:
                # Seek to beginning if it's a file stream
                if hasattr(file_obj_or_path, 'seek'):
                    file_obj_or_path.seek(0)
                reader = pypdf.PdfReader(file_obj_or_path)
                for page in reader.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"

        elif ext == '.docx':
            if isinstance(file_obj_or_path, (str, os.PathLike)):
                doc = docx.Document(file_obj_or_path)
            else:
                if hasattr(file_obj_or_path, 'seek'):
                    file_obj_or_path.seek(0)
                doc = docx.Document(file_obj_or_path)

            paragraphs = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text:
                            paragraphs.append(cell.text)
            text = "\n".join(paragraphs)

        elif ext in ['.txt', '.md']:
            if isinstance(file_obj_or_path, (str, os.PathLike)):
                with open(file_obj_or_path, 'r', encoding='utf-8', errors='ignore') as f:
                    text = f.read()
            else:
                if hasattr(file_obj_or_path, 'seek'):
                    file_obj_or_path.seek(0)
                text = file_obj_or_path.read().decode('utf-8', errors='ignore')

    except Exception as e:
        print(f"Error parsing resume file {filename}: {e}")
        return ""

    return text.strip()
